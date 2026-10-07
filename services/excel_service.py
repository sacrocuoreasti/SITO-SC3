import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, date
from io import BytesIO
from models import db, Persona, NucleoFamiliare, Attivita, Iscrizione
from services.cf_validator import normalizza_cf, valida_cf, estrai_dati_cf

EXCEL_HEADERS = [
    "Codice Fiscale*",
    "Cognome*",
    "Nome*",
    "Sesso (M/F)",
    "Data di Nascita (AAAA-MM-GG)",
    "Luogo di Nascita",
    "Indirizzo Residenza",
    "Comune Residenza",
    "CAP",
    "Telefono",
    "Email",
    "CF Capofamiglia",
    "Nome Famiglia",
    "Ruolo Famiglia",
    "Intolleranze Alimentari",
    "Allergie e Note Mediche",
    "Attivita Iscrizione (Opzionale)"
]

def genera_template_excel() -> BytesIO:
    """
    Genera un file Excel formattato con intestazioni e righe d'esempio per il caricamento parrocchiale.
    Include styling con palette del Sacro Cuore (bordeaux/crimson) e bordi chiari.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anagrafica Sacro Cuore"

    header_fill = PatternFill(start_color="8B1E1E", end_color="8B1E1E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for col_idx, header in enumerate(EXCEL_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    sample_data = [
        [
            "RSSMRA80A01A479L", "Rossi", "Mario", "M", "1980-01-01", "Asti",
            "Via Pier Santi Mattarella 10", "Asti", "14100", "3331234567", "mario.rossi@example.com",
            "RSSMRA80A01A479L", "Famiglia Rossi", "Capofamiglia", "Nessuna", "Nessuna", ""
        ],
        [
            "RSSGNN12B10A479X", "Rossi", "Giovanni", "M", "2012-02-10", "Asti",
            "Via Pier Santi Mattarella 10", "Asti", "14100", "3331234567", "",
            "RSSMRA80A01A479L", "Famiglia Rossi", "Figlio/a", "Celiachia (senza glutine)", "Allergia alle arachidi", "Estate Ragazzi 2026"
        ],
        [
            "BNCFNC85C45A479W", "Bianchi", "Francesca", "F", "1985-03-05", "Asti",
            "Corso Dante 50", "Asti", "14100", "3409876543", "francesca.b@example.com",
            "BNCFNC85C45A479W", "Famiglia Bianchi", "Capofamiglia", "Nessuna", "Nessuna", ""
        ],
        [
            "BNCLRT14D15A479T", "Bianchi", "Lorenzo", "M", "2014-04-15", "Asti",
            "Corso Dante 50", "Asti", "14100", "3409876543", "",
            "BNCFNC85C45A479W", "Famiglia Bianchi", "Figlio/a", "Intolleranza al lattosio", "Punture d'api", "Catechismo - 1ª Comunione"
        ]
    ]

    example_fill = PatternFill(start_color="FDFBF9", end_color="FDFBF9", fill_type="solid")
    for row_idx, row_values in enumerate(sample_data, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = example_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def esporta_anagrafica_excel(persone_list) -> BytesIO:
    """
    Esporta l'elenco delle persone con i dettagli del nucleo familiare in un file Excel.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anagrafica Sacro Cuore"

    header_fill = PatternFill(start_color="8B1E1E", end_color="8B1E1E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = [
        "Codice Fiscale", "Cognome", "Nome", "Sesso", "Età", "Data Nascita",
        "Luogo Nascita", "Indirizzo", "Comune", "CAP", "Telefono", "Email",
        "Famiglia", "Ruolo Famiglia", "Intolleranze Alimentari", "Allergie e Note Mediche"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, p in enumerate(persone_list, start=2):
        ws.cell(row=row_idx, column=1, value=p.codice_fiscale)
        ws.cell(row=row_idx, column=2, value=p.cognome)
        ws.cell(row=row_idx, column=3, value=p.nome)
        ws.cell(row=row_idx, column=4, value=p.sesso or "")
        ws.cell(row=row_idx, column=5, value=p.eta if p.eta is not None else "")
        ws.cell(row=row_idx, column=6, value=p.data_nascita.strftime('%d/%m/%Y') if p.data_nascita else "")
        ws.cell(row=row_idx, column=7, value=p.luogo_nascita or "")
        ws.cell(row=row_idx, column=8, value=p.indirizzo_residenza or "")
        ws.cell(row=row_idx, column=9, value=p.comune_residenza or "")
        ws.cell(row=row_idx, column=10, value=p.cap_residenza or "")
        ws.cell(row=row_idx, column=11, value=p.telefono or "")
        ws.cell(row=row_idx, column=12, value=p.email or "")
        ws.cell(row=row_idx, column=13, value=p.nucleo.nome_famiglia if p.nucleo else "")
        ws.cell(row=row_idx, column=14, value=p.ruolo_famiglia or "")
        ws.cell(row=row_idx, column=15, value=p.intolleranze_alimentari or "")
        ws.cell(row=row_idx, column=16, value=p.allergie or "")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def esporta_iscritti_attivita_excel(attivita) -> BytesIO:
    """
    Esporta gli iscritti di un'attività con colonne per quote e note sanitarie/allergie.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Iscritti - {attivita.titolo[:25]}"

    header_fill = PatternFill(start_color="8B1E1E", end_color="8B1E1E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = [
        "Codice Fiscale", "Cognome", "Nome", "Età", "Data Nascita",
        "Famiglia", "Telefono Genitore/Contatto", "Email",
        "Squadra", "Stato Iscrizione", "Stato Pagamento", "Quota Dovuta (€)", "Quota Pagata (€)",
        "Note / Deleghe", "Intolleranze Alimentari", "Allergie e Note Mediche"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, isc in enumerate(attivita.iscrizioni, start=2):
        p = isc.partecipante
        telefono = p.telefono
        if not telefono and p.nucleo and p.nucleo.capofamiglia:
            telefono = p.nucleo.capofamiglia.telefono

        ws.cell(row=row_idx, column=1, value=p.codice_fiscale)
        ws.cell(row=row_idx, column=2, value=p.cognome)
        ws.cell(row=row_idx, column=3, value=p.nome)
        ws.cell(row=row_idx, column=4, value=p.eta if p.eta is not None else "")
        ws.cell(row=row_idx, column=5, value=p.data_nascita.strftime('%d/%m/%Y') if p.data_nascita else "")
        ws.cell(row=row_idx, column=6, value=p.nucleo.nome_famiglia if p.nucleo else "")
        ws.cell(row=row_idx, column=7, value=telefono or "")
        ws.cell(row=row_idx, column=8, value=p.email or "")
        ws.cell(row=row_idx, column=9, value=isc.squadra or "")
        ws.cell(row=row_idx, column=10, value=isc.stato.capitalize())
        ws.cell(row=row_idx, column=11, value=isc.stato_pagamento.capitalize())
        ws.cell(row=row_idx, column=12, value=isc.importo_dovuto)
        ws.cell(row=row_idx, column=13, value=isc.importo_pagato)
        ws.cell(row=row_idx, column=14, value=isc.note_iscrizione or "")
        ws.cell(row=row_idx, column=15, value=p.intolleranze_alimentari or "")
        ws.cell(row=row_idx, column=16, value=p.allergie or "")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def parse_date(date_val):
    if not date_val:
        return None
    if isinstance(date_val, (date, datetime)):
        return date_val.date() if isinstance(date_val, datetime) else date_val
    date_str = str(date_val).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(date_str, fmt).date()
        except ValueError:
            pass
    return None


def importa_da_excel(file_stream_or_path) -> dict:
    """
    Importa anagrafiche, famiglie e iscrizioni da un file Excel (o BytesIO).
    Esegue l'upsert basandosi sul Codice Fiscale come chiave primaria universale.
    Supporta separatamente intolleranze alimentari e allergie.
    """
    wb = openpyxl.load_workbook(file_stream_or_path, data_only=True)
    ws = wb.active

    risultati = {
        'success': True,
        'totale_righe': 0,
        'persone_create': 0,
        'persone_aggiornate': 0,
        'nuclei_creati': 0,
        'iscrizioni_create': 0,
        'errori': []
    }

    header_row = [cell.value for cell in ws[1]]
    header_map = {}
    for idx, val in enumerate(header_row):
        if val:
            clean_key = str(val).lower().replace('*', '').strip()
            header_map[clean_key] = idx

    def get_val(row_cells, *keywords):
        for kw in keywords:
            for clean_k, col_idx in header_map.items():
                if kw == clean_k:
                    val = row_cells[col_idx].value
                    return str(val).strip() if val is not None else ""
        for kw in keywords:
            for clean_k, col_idx in header_map.items():
                if kw == "nome" and ("cognome" in clean_k or "famiglia" in clean_k):
                    continue
                if kw in ["codice fiscale", "cf"] and "capo" in clean_k:
                    continue
                if kw in clean_k:
                    val = row_cells[col_idx].value
                    return str(val).strip() if val is not None else ""
        return ""

    rows = list(ws.iter_rows(min_row=2))
    risultati['totale_righe'] = len(rows)

    for row_idx, row in enumerate(rows, start=2):
        if all(c.value is None for c in row):
            continue

        cf_raw = get_val(row, "codice fiscale", "cf")
        cf = normalizza_cf(cf_raw)
        
        if not cf:
            risultati['errori'].append(f"Riga {row_idx}: Codice Fiscale assente. Riga ignorata.")
            continue

        nome = get_val(row, "nome")
        cognome = get_val(row, "cognome")
        sesso = get_val(row, "sesso").upper()
        if sesso and sesso[0] in ['M', 'F']:
            sesso = sesso[0]
        else:
            sesso = None

        data_nascita_raw = None
        for k, col_i in header_map.items():
            if "nascita" in k and "data" in k:
                data_nascita_raw = row[col_i].value
                break
        data_nascita = parse_date(data_nascita_raw)

        if not data_nascita or not sesso:
            dati_cf = estrai_dati_cf(cf)
            if not data_nascita and 'data_nascita' in dati_cf:
                data_nascita = dati_cf['data_nascita']
            if not sesso and 'sesso' in dati_cf:
                sesso = dati_cf['sesso']

        luogo_nascita = get_val(row, "luogo")
        indirizzo = get_val(row, "indirizzo")
        comune = get_val(row, "comune") or "Asti"
        cap = get_val(row, "cap") or "14100"
        telefono = get_val(row, "telefono", "cellulare", "tel")
        email = get_val(row, "email", "mail")
        
        intolleranze = get_val(row, "intolleranz", "alimentar", "celiach")
        allergie = get_val(row, "allergi", "medic", "farmac")
        note_generiche = get_val(row, "note sanitarie", "sanitarie")
        if note_generiche and not allergie and not intolleranze:
            allergie = note_generiche

        cf_capo = normalizza_cf(get_val(row, "capofamiglia", "cf capo"))
        nome_famiglia = get_val(row, "nome famiglia", "famiglia")
        ruolo_famiglia = get_val(row, "ruolo") or ("Capofamiglia" if (cf == cf_capo or "capo" in get_val(row, "ruolo").lower()) else "Figlio/a")
        attivita_richiesta = get_val(row, "attivita", "attività")

        persona = db.session.get(Persona, cf)
        if persona:
            if nome: persona.nome = nome
            if cognome: persona.cognome = cognome
            if sesso: persona.sesso = sesso
            if data_nascita: persona.data_nascita = data_nascita
            if luogo_nascita: persona.luogo_nascita = luogo_nascita
            if indirizzo: persona.indirizzo_residenza = indirizzo
            if comune: persona.comune_residenza = comune
            if cap: persona.cap_residenza = cap
            if telefono: persona.telefono = telefono
            if email: persona.email = email
            if intolleranze: persona.intolleranze_alimentari = intolleranze
            if allergie: persona.allergie = allergie
            if ruolo_famiglia: persona.ruolo_famiglia = ruolo_famiglia
            risultati['persone_aggiornate'] += 1
        else:
            if not nome or not cognome:
                risultati['errori'].append(f"Riga {row_idx} ({cf}): Nome e Cognome obbligatori per nuovo inserimento.")
                continue

            persona = Persona(
                codice_fiscale=cf,
                nome=nome,
                cognome=cognome,
                sesso=sesso,
                data_nascita=data_nascita,
                luogo_nascita=luogo_nascita,
                indirizzo_residenza=indirizzo,
                comune_residenza=comune,
                cap_residenza=cap,
                telefono=telefono,
                email=email,
                intolleranze_alimentari=intolleranze,
                allergie=allergie,
                ruolo_famiglia=ruolo_famiglia
            )
            db.session.add(persona)
            risultati['persone_create'] += 1

        db.session.flush()

        # Gestione Nucleo Familiare
        capo_ref = cf_capo if cf_capo else (cf if ruolo_famiglia.lower() == 'capofamiglia' else None)
        if capo_ref:
            nucleo = NucleoFamiliare.query.filter_by(codice_fiscale_capofamiglia=capo_ref).first()
            if not nucleo:
                nome_nuc = nome_famiglia or f"Famiglia {cognome}"
                nucleo = NucleoFamiliare(
                    nome_famiglia=nome_nuc,
                    codice_fiscale_capofamiglia=capo_ref,
                    telefono_principale=telefono,
                    indirizzo=indirizzo,
                    citta=comune
                )
                db.session.add(nucleo)
                db.session.flush()
                risultati['nuclei_creati'] += 1
            
            persona.nucleo_id = nucleo.id

        # Gestione iscrizione rapida da Excel se specificata
        if attivita_richiesta:
            attivita = Attivita.query.filter(Attivita.titolo.ilike(f"%{attivita_richiesta}%")).first()
            if attivita:
                iscrizione_esistente = Iscrizione.query.filter_by(
                    attivita_id=attivita.id,
                    codice_fiscale_partecipante=cf
                ).first()
                if not iscrizione_esistente:
                    nuova_iscrizione = Iscrizione(
                        attivita_id=attivita.id,
                        codice_fiscale_partecipante=cf,
                        stato='confermata',
                        stato_pagamento='da_pagare',
                        importo_dovuto=attivita.quota_iscrizione
                    )
                    db.session.add(nuova_iscrizione)
                    risultati['iscrizioni_create'] += 1

    db.session.commit()
    return risultati
